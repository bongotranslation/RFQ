from __future__ import annotations
from typing import Dict, Any, List
import pathlib
import re
import io
from collections import defaultdict
from docx import Document
from docx.shared import Inches
from PIL import Image

from app.analysis_config import analysis_config

def _count_words(text: str) -> int:
    return len([w for w in text.split() if w.strip()])

def _count_words_by_language(text: str) -> Dict[str, int]:
    """
    Подсчет слов по языкам на основе простой эвристики.
    """
    if not text.strip():
        return {}
    
    words_by_lang = defaultdict(int)
    words = [w for w in text.split() if w.strip()]
    
    # Простая эвристика определения языка по символам
    for word in words:
        # Проверяем на кириллицу (русский)
        if re.search(r'[а-яё]', word.lower()):
            words_by_lang['ru'] += 1
        # Проверяем на латиницу (английский/европейские)
        elif re.search(r'[a-z]', word.lower()):
            words_by_lang['en'] += 1
        # Другие символы - неопределенный язык
        else:
            words_by_lang['other'] += 1
    
    return dict(words_by_lang)

def _is_significant_image(image_part) -> bool:
    """
    Проверяет, является ли изображение значимым (не мелкая иконка/декорация).
    """
    try:
        # Получаем размеры изображения
        image_data = image_part.blob
        img = Image.open(io.BytesIO(image_data))
        width, height = img.size
        
        # Фильтруем мелкие изображения (менее 50x50 пикселей)
        min_size = 50
        if width < min_size or height < min_size:
            return False
            
        # Фильтруем слишком вытянутые изображения (вероятно, линии/разделители)
        aspect_ratio = max(width, height) / min(width, height)
        if aspect_ratio > 10:  # соотношение сторон больше 10:1
            return False
            
        return True
    except Exception:
        return False

def analyze_docx_basic(docx_path: str) -> Dict[str, Any]:
    """
    Анализ DOCX документа с подсчетом текста и изображений:
      - pages_total (приблизительно, основано на количестве текста)
      - by_page: [{p, words, images, needs_ocr}]
      - thresholds
    """
    path = pathlib.Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {path}")

    try:
        doc = Document(str(path))
        
        # Собираем весь текст из параграфов
        all_text = ""
        for paragraph in doc.paragraphs:
            all_text += paragraph.text + "\n"
        
        # Добавляем текст из таблиц
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_text += cell.text + " "
            all_text += "\n"
        
        # Подсчет изображений
        images_count = 0
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                if _is_significant_image(rel.target_part):
                    images_count += 1
        
        # Подсчет слов
        words = _count_words(all_text)
        words_by_language = _count_words_by_language(all_text)
        
        # Приблизительная оценка страниц (500 слов на страницу)
        words_per_page = 500
        estimated_pages = max(1, (words + words_per_page - 1) // words_per_page)
        
        # Распределяем слова и изображения по страницам
        by_page = []
        words_per_estimated_page = words // estimated_pages if estimated_pages > 0 else words
        remaining_words = words
        remaining_images = images_count
        
        for i in range(estimated_pages):
            # Последняя страница получает все оставшиеся слова
            page_words = remaining_words if i == estimated_pages - 1 else words_per_estimated_page
            remaining_words -= page_words
            
            # Распределяем изображения равномерно
            page_images = remaining_images // (estimated_pages - i) if estimated_pages - i > 0 else 0
            remaining_images -= page_images
            
            # Распределяем слова по языкам пропорционально
            page_words_by_lang = {}
            if words > 0:
                for lang, count in words_by_language.items():
                    page_words_by_lang[lang] = int(count * page_words / words)
            
            # OCR нужен если мало текста и есть изображения
            needs_ocr = page_words <= analysis_config.NEEDS_OCR_WORDS_THR and page_images > 0
            
            by_page.append({
                "p": i + 1,
                "words": page_words,
                "words_by_language": page_words_by_lang,
                "images": page_images,
                "needs_ocr": needs_ocr
            })

        return {
            "pages_total": estimated_pages,
            "by_page": by_page,
            "thresholds": {
                "needs_ocr_words_thr": analysis_config.NEEDS_OCR_WORDS_THR,
            },
        }
        
    except Exception as e:
        raise Exception(f"Error analyzing DOCX: {e}")

def update_docx_stats_with_ocr(docx_stats: Dict[str, Any], ocr_results: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Обновляет статистику DOCX с учетом результатов OCR.
    Добавляет OCR слова к исходным словам страниц.
    """
    if not ocr_results or not docx_stats.get("by_page"):
        return docx_stats
    
    # Создаем словарь OCR результатов по номерам страниц
    ocr_by_page = {}
    for ocr_result in ocr_results:
        if isinstance(ocr_result, dict) and "p" in ocr_result:
            page_num = ocr_result["p"]
            ocr_words = ocr_result.get("ocr_words", 0)
            ocr_words_by_language = ocr_result.get("ocr_words_by_language", {})
            ocr_by_page[page_num] = {
                "words": ocr_words,
                "words_by_language": ocr_words_by_language
            }
    
    # Обновляем статистику страниц
    updated_by_page = []
    for page_info in docx_stats["by_page"]:
        page_num = page_info["p"]
        original_words = page_info["words"]
        original_words_by_lang = page_info.get("words_by_language", {})
        
        ocr_data = ocr_by_page.get(page_num, {"words": 0, "words_by_language": {}})
        ocr_words = ocr_data["words"]
        ocr_words_by_lang = ocr_data["words_by_language"]
        
        # Объединяем слова по языкам
        combined_words_by_lang = defaultdict(int)
        for lang, count in original_words_by_lang.items():
            combined_words_by_lang[lang] += count
        for lang, count in ocr_words_by_lang.items():
            combined_words_by_lang[lang] += count
        
        updated_page = page_info.copy()
        updated_page["words"] = original_words + ocr_words
        updated_page["words_by_language"] = dict(combined_words_by_lang)
        updated_page["ocr_applied"] = ocr_words > 0
        
        # Обновляем needs_ocr - если OCR уже применен, то больше не нужен
        if ocr_words > 0:
            updated_page["needs_ocr"] = False
            
        updated_by_page.append(updated_page)
    
    # Обновляем статистику
    updated_stats = docx_stats.copy()
    updated_stats["by_page"] = updated_by_page
    
    return updated_stats
